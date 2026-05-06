"""Document text extraction: txt, md, csv, pdf, docx.

Pure functions. Blocking I/O — call from handlers via `asyncio.to_thread`.
Format resolved by file extension first, mime-type second: Telegram often
sends `application/octet-stream`, so trusting the extension reduces false
"unsupported" rejections.
"""

from __future__ import annotations

import csv
import os
from typing import Optional


class DocExtractError(RuntimeError):
    """Base for all document-extraction failures."""


class UnsupportedDocumentError(DocExtractError):
    """File type isn't on the supported list."""


class EmptyDocumentError(DocExtractError):
    """Parser ran but produced no text — e.g. scan-only or encrypted PDF."""


_TXT_EXTS = {".txt"}
_MD_EXTS = {".md", ".markdown"}
_CSV_EXTS = {".csv"}
_PDF_EXTS = {".pdf"}
_DOCX_EXTS = {".docx"}

_TXT_MIMES = {"text/plain"}
_MD_MIMES = {"text/markdown", "text/x-markdown"}
_CSV_MIMES = {"text/csv"}
_PDF_MIMES = {"application/pdf"}
_DOCX_MIMES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
}

_SNIFF_BYTES = 4096


def _resolve_kind(mime: Optional[str], file_name: Optional[str]) -> str:
    ext = ""
    if file_name:
        ext = os.path.splitext(file_name)[1].lower()

    if ext in _TXT_EXTS:
        return "txt"
    if ext in _MD_EXTS:
        return "md"
    if ext in _CSV_EXTS:
        return "csv"
    if ext in _PDF_EXTS:
        return "pdf"
    if ext in _DOCX_EXTS:
        return "docx"

    mime_l = (mime or "").lower()
    if mime_l in _TXT_MIMES:
        return "txt"
    if mime_l in _MD_MIMES:
        return "md"
    if mime_l in _CSV_MIMES:
        return "csv"
    if mime_l in _PDF_MIMES:
        return "pdf"
    if mime_l in _DOCX_MIMES:
        return "docx"

    return "unknown"


def extract(path: str, mime_type: Optional[str], file_name: Optional[str]) -> str:
    """Extract plain text from `path` using the format inferred from name/mime.

    Raises:
        UnsupportedDocumentError: format not supported and bytes don't look like text.
        EmptyDocumentError: parser succeeded but text is empty (e.g. scanned PDF).
        DocExtractError: any other parser failure.
    """
    kind = _resolve_kind(mime_type, file_name)

    if kind in ("txt", "md"):
        return _extract_text_file(path)
    if kind == "csv":
        return _extract_csv(path)
    if kind == "pdf":
        return _extract_pdf(path)
    if kind == "docx":
        return _extract_docx(path)

    return _extract_unknown_fallback(path)


def _extract_text_file(path: str) -> str:
    with open(path, "rb") as f:
        data = f.read()
    return data.decode("utf-8-sig", errors="replace")


def _extract_csv(path: str) -> str:
    rows: list[str] = []
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
        reader = csv.reader(f)
        for row in reader:
            cells = [c.strip() for c in row[:50]]
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _extract_pdf(path: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocExtractError(f"pypdf not available: {exc}") from exc

    try:
        reader = PdfReader(path)
    except Exception as exc:
        raise DocExtractError(f"pdf read failed: {exc}") from exc

    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    text = "\n\n".join(p for p in pages if p).strip()
    if not text:
        raise EmptyDocumentError(
            "no extractable text — looks like a scan or encrypted PDF"
        )
    return text


def _extract_docx(path: str) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocExtractError(f"python-docx not available: {exc}") from exc

    try:
        doc = Document(path)
    except Exception as exc:
        raise DocExtractError(f"docx read failed: {exc}") from exc

    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text:
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append("\t".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise EmptyDocumentError("docx contains no text")
    return text


def _extract_unknown_fallback(path: str) -> str:
    with open(path, "rb") as f:
        head = f.read(_SNIFF_BYTES)
    try:
        head.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise UnsupportedDocumentError(
            "поддерживаются: txt, md, csv, pdf, docx"
        ) from exc
    return _extract_text_file(path)


__all__ = [
    "DocExtractError",
    "UnsupportedDocumentError",
    "EmptyDocumentError",
    "extract",
]
